import io
import json
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook

from assets import summary_export


def program(pid, school, name, fit, *, facts=None, advisors=(), deadline_status=None):
    record = {
        "program_id": pid,
        "region": "US",
        "school": school,
        "program": name,
        "facts": {"fitScore": fit, **(facts or {})},
        "pis": {"pis": list(advisors)},
        "provenance": {"checks": []},
    }
    if deadline_status:
        record["provenance"]["checks"].append(
            {"claimPath": "facts.deadlineAndTests", "status": deadline_status}
        )
    return record


RESEARCH = {
    "Example University": [
        program(
            "program_a", "Example University", "Biology PhD", 8.5,
            facts={
                "applicationDeadline": "2026-11-20",
                "programUrl": "https://example.edu/bio",
                "applicationPortalUrl": "https://apply.example.edu",
            },
            advisors=[
                {"advisor_id": "adv_1", "name": "Alex Li", "url": "https://lab.example.edu"},
                {"advisor_id": "adv_2", "name": "Sam Hidden", "url": ""},
            ],
            deadline_status="confirmed",
        ),
        program(
            "program_b", "Example University", "Neuroscience PhD", 6,
            facts={"deadlineAndTests": "Apply by Jan 5, 2027; GRE optional"},
            advisors=[{"advisor_id": "adv_3", "name": "Alex Li", "url": "https://lab.example.edu"}],
        ),
    ],
    "Other Institute": [
        program("program_c", "Other Institute", "=Hidden Program", 3),
        program("program_d", "Other Institute", "+Systems Biology", None),
    ],
}
HIDDEN = {"programs": ["program_c"], "pis": ["adv_2"]}


class SummaryTablesTests(unittest.TestCase):
    def setUp(self):
        self.tables = summary_export.summary_tables(RESEARCH, hidden=HIDDEN)

    def test_rows_are_grouped_and_tracker_fields_derived(self):
        programs = self.tables["programs"]
        self.assertEqual([p["program_id"] for p in programs], ["program_a", "program_b", "program_d"])
        a, b, d = programs
        self.assertEqual((a["deadline"], a["match_level"]), ("2026-11-20", "High"))
        # Parsed from free text and not independently confirmed.
        self.assertEqual((b["deadline"], b["match_level"]), ("2027-01-05 (TBC)", "Medium"))
        self.assertEqual((d["deadline"], d["match_level"], d["degree_type"]), ("TBC", "", "PhD"))

    def test_hidden_advisors_excluded_and_people_deduplicated_per_school(self):
        names = [a["name"] for a in self.tables["advisors"]]
        self.assertEqual(names, ["Alex Li"])
        self.assertEqual(self.tables["advisors"][0]["notes"], "")

    def test_program_filter_and_notes_opt_in(self):
        tables = summary_export.summary_tables(
            RESEARCH, program_ids=["program_b"], notes={"adv_3": "met at conference"}
        )
        self.assertEqual([p["program_id"] for p in tables["programs"]], ["program_b"])
        self.assertEqual(tables["advisors"][0]["notes"], "met at conference")


class SummaryWritersTests(unittest.TestCase):
    def setUp(self):
        self.tables = summary_export.summary_tables(RESEARCH, hidden=HIDDEN)

    def test_xlsx_follows_tracker_layout(self):
        book = load_workbook(io.BytesIO(summary_export.build_summary_xlsx(self.tables)))
        self.assertEqual(
            book.sheetnames, ["Application overall", "Application status", "Reach out status"]
        )
        overview = book["Application overall"]
        self.assertEqual(overview["A1"].value, "PhD Programs")
        self.assertEqual(
            [c.value for c in overview[2]],
            [name for name, _ in summary_export.OVERVIEW_COLUMNS],
        )
        merged = {str(r) for r in overview.merged_cells.ranges}
        self.assertIn("D3:D4", merged)  # Example University spans its two programs
        self.assertEqual(overview["E3"].fill.fgColor.rgb, "FFF4CCCC")
        self.assertTrue(overview["E5"].value.startswith("'+"))
        status = book["Application status"]
        self.assertEqual(str(status.data_validations.dataValidation[0].sqref), "H3:H5")
        for sheet in book.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    self.assertNotEqual(cell.data_type, "f")

    def test_html_is_google_docs_importable_and_escaped(self):
        page = summary_export.build_summary_html(self.tables)
        self.assertEqual(page.count("<table"), 3)
        self.assertIn("rowspan='2'", page)
        self.assertIn("href='https://apply.example.edu'", page)
        self.assertNotIn("Hidden Program", page)
        self.assertNotIn("Sam Hidden", page)

    def test_cli_writes_both_formats(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp)
            (out / "_research_data.json").write_text(json.dumps(RESEARCH))
            (out / "_pi_hidden.json").write_text(json.dumps(HIDDEN))
            self.assertEqual(summary_export.main([str(out), str(out / "t.xlsx")]), 0)
            self.assertEqual(summary_export.main([str(out), str(out / "t.html")]), 0)
            self.assertEqual(len(load_workbook(out / "t.xlsx")["Application overall"]["A"]), 7)
            self.assertIn("<table", (out / "t.html").read_text())
            self.assertEqual(summary_export.main([str(out / "missing"), str(out / "x.xlsx")]), 2)


if __name__ == "__main__":
    unittest.main()
