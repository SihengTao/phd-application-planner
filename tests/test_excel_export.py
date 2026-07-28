import io
import json
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook
import pandas as pd

from assets import excel_export


def workbook_from(payload):
    return load_workbook(io.BytesIO(payload), data_only=False)


class ExcelExportTests(unittest.TestCase):
    def setUp(self):
        self.programs = pd.DataFrame(
            [
                {
                    "program_id": "program_1",
                    "school": "Example",
                    "program": "History PhD",
                    "fit_rationale": "=HYPERLINK(\"https://bad.example\")",
                }
            ]
        )
        self.advisors = pd.DataFrame(
            [
                {
                    "advisor_id": "advisor_1",
                    "program_id": "program_1",
                    "name": "Alex Li",
                    "my_private_note": "@private",
                    "research": "+SUM(1,2)",
                }
            ]
        )

    def build(self, *, include_notes=False):
        return excel_export.build_excel_bytes(
            self.programs,
            self.advisors,
            pd.DataFrame([{"name": "Red Pepper", "rating": 4.5}]),
            pd.DataFrame([{"url": "https://example.edu"}]),
            pd.DataFrame([{"severity": "warning", "message": "-review"}]),
            {
                "field": "History",
                "source_directory": "/Users/private/run",
                "includes_private_notes": include_notes,
            },
            priority_df=pd.DataFrame([{"program_id": "program_1", "score": 0.9}]),
            include_notes=include_notes,
        )

    def test_workbook_structure_formatting_and_formula_neutralization(self):
        workbook = workbook_from(self.build())
        self.assertEqual(
            workbook.sheetnames,
            [
                "Programs",
                "Advisors",
                "Restaurants",
                "Sources",
                "Quality",
                "Priority",
                "Metadata",
            ],
        )
        self.assertEqual(workbook["Programs"].freeze_panes, "A2")
        self.assertIsNotNone(workbook["Programs"].auto_filter.ref)
        self.assertTrue(
            workbook["Programs"]["D2"].value.startswith("'="),
            workbook["Programs"]["D2"].value,
        )
        self.assertTrue(
            workbook["Advisors"]["D2"].value.startswith("'+"),
            workbook["Advisors"]["D2"].value,
        )
        for worksheet in workbook.worksheets:
            for row in worksheet.iter_rows():
                for cell in row:
                    self.assertNotEqual(cell.data_type, "f")
                    self.assertIsNone(cell.hyperlink)
        metadata_values = [
            cell.value for cell in workbook["Metadata"]["B"] if cell.value
        ]
        self.assertFalse(any("/Users/private" in str(value) for value in metadata_values))

    def test_private_notes_are_opt_in(self):
        default_book = workbook_from(self.build(include_notes=False))
        included_book = workbook_from(self.build(include_notes=True))
        default_headers = [
            cell.value for cell in default_book["Advisors"][1]
        ]
        included_headers = [
            cell.value for cell in included_book["Advisors"][1]
        ]
        self.assertNotIn("my_private_note", default_headers)
        self.assertIn("my_private_note", included_headers)
        note_column = included_headers.index("my_private_note") + 1
        self.assertEqual(
            included_book["Advisors"].cell(2, note_column).value,
            "'@private",
        )

    def test_cli_exports_built_data_and_joins_notes_by_id_only(self):
        record = {
            "program_id": "program_1",
            "region": "US",
            "school": "Example",
            "program": "History PhD",
            "facts": {"applicationPolicy": "single_only"},
            "pis": {
                "pis": [
                    {
                        "advisor_id": "advisor_1",
                        "name": "Alex Li",
                        "research": "Archives",
                    },
                    {
                        "advisor_id": "advisor_2",
                        "name": "Alex Li",
                        "research": "Methods",
                    },
                ]
            },
            "out": {},
            "nearbyFood": {"restaurants": [], "sources": []},
            "provenance": {"sources": []},
        }
        row = {
            "program_id": "program_1",
            "region": "US",
            "school": "Example",
            "program": "History PhD",
        }
        with tempfile.TemporaryDirectory() as temporary:
            out_dir = Path(temporary)
            (out_dir / "_research_data.json").write_text(
                json.dumps({"Example": [record]}), encoding="utf-8"
            )
            (out_dir / "_rows.json").write_text(
                json.dumps([row]), encoding="utf-8"
            )
            (out_dir / "_pi_notes.json").write_text(
                json.dumps(
                    {
                        "advisor_1": "ID-specific note",
                        "Alex Li": "ambiguous legacy note",
                    }
                ),
                encoding="utf-8",
            )
            default_path = out_dir / "default.xlsx"
            private_path = out_dir / "private.xlsx"
            self.assertEqual(
                excel_export.main([str(out_dir), str(default_path)]), 0
            )
            self.assertEqual(
                excel_export.main(
                    [str(out_dir), str(private_path), "--include-notes"]
                ),
                0,
            )
            default_book = load_workbook(default_path)
            private_book = load_workbook(private_path)

        default_headers = [
            cell.value for cell in default_book["Advisors"][1]
        ]
        private_headers = [
            cell.value for cell in private_book["Advisors"][1]
        ]
        self.assertNotIn("my_private_note", default_headers)
        note_column = private_headers.index("my_private_note") + 1
        self.assertEqual(private_book["Advisors"].cell(2, note_column).value, "ID-specific note")
        self.assertIsNone(private_book["Advisors"].cell(3, note_column).value)

    def test_cli_preserves_sources_from_all_supported_locations(self):
        shared_source = {
            "id": "shared-source",
            "url": "https://example.edu/shared",
        }
        record = {
            "program_id": "program_1",
            "school": "Example",
            "program": "History PhD",
            "facts": {
                "sourceRecords": [{"id": "facts-record", "url": "https://example.edu/facts"}],
                "sources": ["https://example.edu/legacy-facts", shared_source],
            },
            "pis": {
                "pis": [
                    {
                        "advisor_id": "advisor_1",
                        "name": "Alex Li",
                        "sources": [
                            {"id": "advisor-source", "url": "https://example.edu/alex"},
                            shared_source,
                        ],
                    }
                ],
                "sources": [{"id": "faculty-source", "url": "https://example.edu/faculty"}],
            },
            "out": {
                "sourceRecords": [{"id": "outcome-source", "url": "https://example.edu/outcomes"}]
            },
            "nearbyFood": {
                "restaurants": [],
                "sources": [{"id": "food-source", "url": "https://example.com/food"}],
            },
            "provenance": {
                "sources": [
                    {"id": "program-source", "url": "https://example.edu/program"},
                    shared_source,
                ],
                "checks": [
                    {"claimPath": "facts.stipendLocal", "status": "corrected"},
                    {"claimPath": "facts.deadlineAndTests", "status": "unresolved"},
                ],
                "conflicts": [
                    {"claimPath": "facts.applicationPolicy", "status": "open"}
                ],
                "researchStatus": "conflict",
            },
            "verification": {
                "status": "conflict",
                "checkedAt": "2026-07-27T18:45:00-04:00",
                "conflict": True,
            },
        }
        with tempfile.TemporaryDirectory() as temporary:
            out_dir = Path(temporary)
            (out_dir / "_research_data.json").write_text(
                json.dumps({"Example": [record]}), encoding="utf-8"
            )
            (out_dir / "_rows.json").write_text(
                json.dumps([{"program_id": "program_1"}]), encoding="utf-8"
            )
            flattened = excel_export._flatten_cli(out_dir, include_notes=False)

        urls = set(flattened["sources"]["url"].tolist())
        self.assertEqual(
            urls,
            {
                "https://example.edu/program",
                "https://example.edu/facts",
                "https://example.edu/legacy-facts",
                "https://example.edu/faculty",
                "https://example.edu/alex",
                "https://example.edu/outcomes",
                "https://example.com/food",
                "https://example.edu/shared",
            },
        )
        shared_scopes = set(
            flattened["sources"].loc[
                flattened["sources"]["url"] == "https://example.edu/shared",
                "scope",
            ]
        )
        self.assertEqual(
            shared_scopes,
            {"program", "facts:sources", "advisor:advisor_1"},
        )
        program = flattened["programs"].iloc[0]
        self.assertEqual(program["verification_status"], "conflict")
        self.assertEqual(program["provenance_research_status"], "conflict")
        self.assertEqual(program["check_count"], 2)
        self.assertEqual(program["conflict_count"], 1)
        self.assertEqual(program["unresolved_check_count"], 1)
        self.assertEqual(program["corrected_check_count"], 1)


if __name__ == "__main__":
    unittest.main()
